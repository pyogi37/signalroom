"""DOCX export of a room's brief, in the working shape of a solution architecture spec."""

from io import BytesIO

from docx import Document

from .models import Room


def _table(document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    document.add_paragraph()


def brief_docx(room: Room) -> BytesIO:
    document = Document()
    document.add_heading(f"Solution brief: {room.organization}", 0)
    document.add_paragraph(
        "Synthetic working material produced by SignalRoom. Every organization, person and number in this document is invented. "
        + ("Approved by the solution engineer." if room.status == "approved" else "Not yet approved; human review required before use.")
    )

    brief = room.brief
    if brief:
        document.add_heading("Snapshot", level=1)
        _table(document, ["Field", "Value"], [
            ["Goal", brief.snapshot.one_line_goal],
            ["Scope", brief.snapshot.scope_summary],
            ["Deployment shape", brief.snapshot.deployment_shape],
            ["Stakeholder roles", ", ".join(brief.snapshot.stakeholder_roles) or "TBC"],
            ["Brief revision", str(brief.revision)],
        ])

    document.add_heading("Use cases", level=1)
    readiness = {item.use_case_id: item for item in (brief.readiness if brief else [])}
    if room.use_cases:
        _table(document, ["Id", "Use case", "Trigger condition", "Expected output", "Readiness", "Evidence"], [
            [case.id, case.name, case.trigger_condition, case.expected_output,
             readiness[case.id].readiness.replace("_", " ") if case.id in readiness else "unknown",
             f"L{case.evidence.line} {case.evidence.speaker}: \"{case.evidence.quote}\""]
            for case in room.use_cases
        ])
    else:
        document.add_paragraph("No use cases were established in the conversation.")

    document.add_heading("Requirements with evidence", level=1)
    for item in room.requirements:
        document.add_heading(f"{item.id}  {item.title}", level=2)
        document.add_paragraph(f"{item.kind.replace('_', ' ').capitalize()}. Confidence {item.confidence}: {item.confidence_reason}")
        document.add_paragraph(item.detail)
        document.add_paragraph(f"L{item.evidence.line} {item.evidence.speaker}: \"{item.evidence.quote}\"")

    if brief:
        document.add_heading("Environment and constraints", level=1)
        for item in brief.constraints:
            origin = f"L{item.line}" if item.line else (item.passage_id or item.source)
            document.add_paragraph(f"{item.constraint} ({item.source.replace('_', ' ')}, {origin})", style="List Bullet")

        document.add_heading("Recommendation", level=1)
        document.add_paragraph(brief.recommendation.approach)
        _table(document, ["Phase", "Purpose", "Duration", "Exit criteria"], [
            [phase.name, phase.purpose, phase.duration, "; ".join(phase.exit_criteria)] for phase in brief.recommendation.phases
        ])
        if brief.recommendation.discussed_not_in_scope:
            document.add_heading("Discussed but not in scope", level=2)
            for item in brief.recommendation.discussed_not_in_scope:
                document.add_paragraph(item, style="List Bullet")

        document.add_heading("Success measures", level=1)
        _table(document, ["Measure", "Baseline", "Owner role"], [
            [item.measure, item.baseline_status.replace("_", " "), item.owner_role] for item in brief.success_measures
        ])

        document.add_heading("Risks", level=1)
        _table(document, ["Risk", "Severity", "Mitigation", "Basis"], [
            [item.title, item.severity, item.mitigation, item.basis] for item in brief.risks
        ])

    document.add_heading("Open items", level=1)
    _table(document, ["Id", "Question", "Owner role", "Status"], [
        [item.id, item.question, item.suggested_owner_role,
         f"answered (L{item.answered_line}): {item.answer}" if item.status == "answered" else "open"]
        for item in room.open_items
    ])

    if room.critique:
        document.add_heading("Critic pass", level=1)
        document.add_paragraph(f"Verdict: {room.critique.verdict.replace('_', ' ')}. {room.critique.summary}")
        for finding in room.critique.findings:
            document.add_paragraph(f"[{finding.severity}] {finding.kind.replace('_', ' ')}: {finding.text}", style="List Bullet")

    document.add_heading("Provenance", level=1)
    metrics = room.metrics or {}
    document.add_paragraph(
        f"Model calls: {len(metrics.get('calls', []))}. Modes: {', '.join(metrics.get('modes', [])) or 'none'}. "
        f"Grounding: {room.grounding.passed} of {room.grounding.proposed} proposed claims passed the quote check, "
        f"{room.grounding.repaired} repaired, {room.grounding.dropped} dropped."
    )
    for decision in room.decisions:
        document.add_paragraph(f"{decision.decided_at}: {decision.kind.replace('_', ' ')}{(' - ' + decision.note) if decision.note else ''}", style="List Bullet")

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
